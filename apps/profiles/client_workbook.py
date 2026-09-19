from __future__ import annotations

import csv
import io
import math
from dataclasses import dataclass
from pathlib import Path

from django.conf import settings


MAX_CLIENT_RAW_ROWS = 100_002


class ClientWorkbookError(ValueError):
    pass


@dataclass(frozen=True)
class ClientWorkbookData:
    pv_base_profile: list[float]
    load_profile: list[float]
    row_count: int
    detected_format: str
    warnings: list[str]


def _empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _read_rows(name: str, raw: bytes) -> tuple[list[list], str, int]:
    """Read the first two columns with a hard raw-row cap for resource safety."""
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".xls":
            import xlrd

            book = xlrd.open_workbook(file_contents=raw, on_demand=True)
            try:
                sheet = book.sheet_by_index(0)
                if sheet.nrows > MAX_CLIENT_RAW_ROWS:
                    raise ClientWorkbookError("The workbook contains too many rows.")
                max_columns = int(sheet.ncols)
                rows = [
                    [sheet.cell_value(r, c) for c in range(min(sheet.ncols, 2))]
                    for r in range(sheet.nrows)
                ]
            finally:
                book.release_resources()
            return rows, "xls", max_columns

        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                max_columns = int(sheet.max_column or 0)
                rows = []
                for index, row in enumerate(
                    sheet.iter_rows(min_col=1, max_col=2, values_only=True),
                    start=1,
                ):
                    if index > MAX_CLIENT_RAW_ROWS:
                        raise ClientWorkbookError("The workbook contains too many rows.")
                    rows.append(list(row))
            finally:
                workbook.close()
            return rows, "xlsx", max_columns

        if suffix in {".csv", ".txt"}:
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("latin-1")
            rows = []
            max_columns = 0
            for index, row in enumerate(csv.reader(io.StringIO(text)), start=1):
                if index > MAX_CLIENT_RAW_ROWS:
                    raise ClientWorkbookError("The workbook contains too many rows.")
                max_columns = max(max_columns, len(row))
                rows.append(list(row[:2]))
            return rows, suffix.lstrip("."), max_columns
    except ClientWorkbookError:
        raise
    except ImportError as exc:
        dependency = "xlrd" if suffix == ".xls" else "openpyxl"
        raise ClientWorkbookError(
            f"{suffix.upper()} support requires the {dependency!r} package. "
            "Install backend requirements first."
        ) from exc
    except Exception as exc:
        raise ClientWorkbookError(f"Could not read {suffix or 'uploaded'} file.") from exc

    raise ClientWorkbookError("Supported client workbook formats are XLS, XLSX, CSV and TXT.")


def _numeric_pair(row) -> tuple[float, float]:
    if len(row) < 2:
        raise ValueError("row has fewer than two columns")
    pv = float(row[0])
    load = float(row[1])
    if not math.isfinite(pv) or not math.isfinite(load):
        raise ValueError("non-finite number")
    return pv, load


def parse_client_workbook(uploaded_file) -> ClientWorkbookData:
    """
    Parse the client workbook convention used by natpaymenttrail.m.

    Column 1 -> legacy/base PV profile
    Column 2 -> household load profile

    A single non-numeric header row is tolerated. Empty rows are ignored and
    additional columns are ignored with a warning.
    """
    max_bytes = settings.PROFILE_UPLOAD_MAX_MB * 1024 * 1024
    if getattr(uploaded_file, "size", 0) and uploaded_file.size > max_bytes:
        raise ClientWorkbookError(
            f"File exceeds the {settings.PROFILE_UPLOAD_MAX_MB} MB upload limit."
        )

    raw = uploaded_file.read()
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)
    if not raw:
        raise ClientWorkbookError("The uploaded file is empty.")

    rows, file_kind, max_columns = _read_rows(uploaded_file.name, raw)
    rows = [row for row in rows if row and any(not _empty(cell) for cell in row)]
    if not rows:
        raise ClientWorkbookError("The workbook contains no data rows.")
    if max_columns < 2:
        raise ClientWorkbookError(
            "The client workbook needs at least two columns: column 1 = PV base profile, column 2 = load."
        )

    warnings: list[str] = []
    if max_columns > 2:
        warnings.append(
            f"The workbook contains {max_columns} columns; only the first two are used by the client algorithm."
        )

    header_removed = False
    start = 0
    try:
        _numeric_pair(rows[0])
    except Exception:
        header_removed = True
        start = 1

    pv_values: list[float] = []
    load_values: list[float] = []
    for row_number, row in enumerate(rows[start:], start=start + 1):
        if len(row) < 2 or (_empty(row[0]) and _empty(row[1])):
            continue
        try:
            pv, load = _numeric_pair(row)
        except Exception as exc:
            raise ClientWorkbookError(
                f"Row {row_number} must contain numeric PV and load values in the first two columns."
            ) from exc
        if pv < 0 or load < 0:
            raise ClientWorkbookError(
                f"Row {row_number} contains a negative PV/load value, which is not supported by the client algorithm."
            )
        pv_values.append(pv)
        load_values.append(load)

    if not pv_values:
        raise ClientWorkbookError("No numeric PV/load rows were found.")
    if len(pv_values) < 2:
        raise ClientWorkbookError("At least two hourly PV/load rows are required.")
    if len(pv_values) > 100_000:
        raise ClientWorkbookError("The workbook contains more than 100,000 rows and is not accepted.")

    if header_removed:
        warnings.append("One header row was detected and ignored.")
    if len(load_values) not in {8760, 8784}:
        warnings.append(
            f"The workbook contains {len(load_values)} rows. The client algorithm treats each row as one hour; "
            "annual-labelled results are only annual when the file represents a complete 8760/8784-hour year."
        )

    return ClientWorkbookData(
        pv_base_profile=pv_values,
        load_profile=load_values,
        row_count=len(load_values),
        detected_format=f"client_two_column_{file_kind}",
        warnings=warnings,
    )
